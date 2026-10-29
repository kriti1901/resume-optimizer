"""Turn an uploaded resume (PDF, DOCX or TXT) into plain text, and strip personal info."""

import io
import re

import pdfplumber
from docx import Document

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d[\d\s().-]{8,}\d)")
URL_RE = re.compile(r"(?:https?://|www\.)\S+|(?:linkedin|github)\.com/\S+", re.I)


def parse_resume(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT file.")
    return normalize(text)


def normalize(text: str) -> str:
    text = text.replace("•", "-").replace("●", "-").replace("", "-")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def strip_pii(text: str) -> str:
    """Remove email, phone numbers and profile URLs before sending text to the LLM."""
    text = EMAIL_RE.sub("[EMAIL]", text)
    text = URL_RE.sub("[URL]", text)
    text = PHONE_RE.sub("[PHONE]", text)
    return text
